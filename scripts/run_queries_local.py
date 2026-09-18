"""
run_queries_local.py
--------------------------------------------------------------------------
Reproduce every result file in results/ from the four raw CSVs, without a
SQL Server connection.

The SQL files in sql/ are written in T-SQL for XómDataset (SQL Server).
This script loads the CSVs into an in-memory SQLite database (attached as
schema `e_commerce`, so `e_commerce.ecom_sales` resolves unchanged) and
registers the few T-SQL scalar functions SQLite lacks (YEAR, MONTH,
DATEDIFF). Only two syntax differences are rewritten mechanically:

    DATEDIFF(day, a, b)   ->  DATEDIFF_DAY(a, b)
    SELECT TOP (n) ...    ->  SELECT ... LIMIT n   (outer query only)

Everything else (CTEs, window functions, NTILE, LAG, FIRST_VALUE, CONCAT,
conditional aggregation) is identical in both dialects, so the T-SQL file
is the single source of truth and the results are reproducible.

Usage:
    python scripts/run_queries_local.py --data data --sql sql --out results

One SQL file per domain (sql/01_finance_pnl.sql ...). Each `-- @output <name>`
marker inside it starts one statement; its result becomes one sheet in
results/<domain>.xlsx (sheet name = the short id, e.g. fin_01b; an INDEX
sheet in every workbook maps ids to full names and descriptions). A
results/_index.md listing every output is generated as well.
"""

import argparse
import datetime as dt
import re
import sqlite3
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------- helpers
def _parse_date(value):
    if value is None:
        return None
    return dt.date.fromisoformat(str(value)[:10])


def register_tsql_functions(con: sqlite3.Connection) -> None:
    con.create_function("YEAR", 1, lambda d: _parse_date(d).year if d else None)
    con.create_function("MONTH", 1, lambda d: _parse_date(d).month if d else None)
    con.create_function("DAY", 1, lambda d: _parse_date(d).day if d else None)
    # T-SQL DATEDIFF(day, start, end) = end - start
    con.create_function(
        "DATEDIFF_DAY", 2,
        lambda a, b: (_parse_date(b) - _parse_date(a)).days if a and b else None,
    )
    con.create_function("LEN", 1, lambda s: len(s) if s is not None else None)


def load_csvs(con: sqlite3.Connection, data_dir: Path) -> None:
    con.execute("ATTACH DATABASE ':memory:' AS e_commerce")
    files = {
        "customer":   "e_commerce_customer.csv",
        "ecom_sales": "e_commerce_ecom_sales.csv",
        "product":    "e_commerce_product.csv",
        "region":     "e_commerce_region.csv",
    }
    for table, fname in files.items():
        path = data_dir / fname
        if not path.exists():
            sys.exit(f"missing data file: {path}")
        df = pd.read_csv(path)
        df.to_sql(f"tmp_{table}", con, index=False, if_exists="replace")
        con.execute(f"CREATE TABLE e_commerce.{table} AS SELECT * FROM main.tmp_{table}")
        con.execute(f"DROP TABLE main.tmp_{table}")
        print(f"  loaded e_commerce.{table:<11} {len(df):>7,} rows")
    # indexes on join keys keep the self-join / RFM queries fast
    con.execute("CREATE INDEX e_commerce.ix_sales_order    ON ecom_sales(order_id)")
    con.execute("CREATE INDEX e_commerce.ix_sales_customer ON ecom_sales(customer_id)")
    con.execute("CREATE INDEX e_commerce.ix_sales_product  ON ecom_sales(product_code)")
    con.execute("CREATE INDEX e_commerce.ix_sales_region   ON ecom_sales(region_code)")


TOP_RE = re.compile(r"SELECT\s+TOP\s*\(?\s*(\d+)\s*\)?", re.IGNORECASE)
DATEDIFF_RE = re.compile(r"DATEDIFF\s*\(\s*day\s*,", re.IGNORECASE)


def to_sqlite(statement: str) -> str:
    sql = DATEDIFF_RE.sub("DATEDIFF_DAY(", statement)
    tops = TOP_RE.findall(sql)
    if len(tops) > 1:
        raise ValueError("TOP is only supported once per statement (outer query)")
    if tops:
        sql = TOP_RE.sub("SELECT", sql, count=1).rstrip().rstrip(";")
        sql = f"{sql}\nLIMIT {tops[0]}"
    return sql


MARKER_RE = re.compile(r"^--\s*@output\s+(\S+)\s*$", re.MULTILINE)
SHEET_ID_RE = re.compile(r"^([a-z]+_\d+[a-z]?)_")


def split_outputs(sql_text: str):
    """Yield (name, description, statement) for every `-- @output name` block.
    The description is the run of `--` comment lines directly under the marker."""
    parts = MARKER_RE.split(sql_text)
    # parts = [preamble, name1, body1, name2, body2, ...]
    for i in range(1, len(parts), 2):
        name, body = parts[i], parts[i + 1].strip()
        desc_lines = []
        for line in body.splitlines():
            if line.startswith("--"):
                desc_lines.append(line[2:].strip())
            else:
                break
        yield name, " ".join(desc_lines), strip_trailing_comments(body)


TRAILING_BLOCK_RE = re.compile(r"/\*.*?\*/\s*$", re.DOTALL)
TRAILING_LINE_RE = re.compile(r"(?:^|\n)--[^\n]*\s*$")


def strip_trailing_comments(body: str) -> str:
    """Drop the next section's header comment that trails the last statement."""
    prev = None
    while prev != body:
        prev = body
        body = TRAILING_BLOCK_RE.sub("", body).rstrip()
        body = TRAILING_LINE_RE.sub("", body).rstrip()
    return body.rstrip(";").strip()


def sheet_id(name: str) -> str:
    m = SHEET_ID_RE.match(name)
    return m.group(1) if m else name[:31]


def tidy(df: pd.DataFrame) -> pd.DataFrame:
    for col in df.columns:
        if df[col].dtype == "float64":
            df[col] = df[col].round(4)
    return df


# ------------------------------------------------------------- xlsx output
FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF")
BODY_FONT = Font(name=FONT)
TITLE_FONT = Font(name=FONT, bold=True, size=13)


def write_sheet(wb: Workbook, title: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(title)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
    for cell in ws[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for col_idx, col in enumerate(df.columns, start=1):
        letter = get_column_letter(col_idx)
        sample = df[col].astype(str).head(200).str.len().max() if len(df) else 0
        ws.column_dimensions[letter].width = max(10, min(45, max(len(str(col)), int(sample or 0)) + 2))
        if pd.api.types.is_float_dtype(df[col]):
            fmt = "0.00" if col.endswith("_pct") or col in ("lift",) else "#,##0.00"
            for cell in ws[letter][1:]:
                cell.number_format = fmt
                cell.font = BODY_FONT
        elif pd.api.types.is_integer_dtype(df[col]):
            for cell in ws[letter][1:]:
                cell.number_format = "#,##0"
                cell.font = BODY_FONT
        else:
            for cell in ws[letter][1:]:
                cell.font = BODY_FONT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_index_sheet(wb: Workbook, domain: str, sql_file: str, rows: list) -> None:
    ws = wb.create_sheet("INDEX", 0)
    ws["A1"] = f"{domain} - query results"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (f"Source: sql/{sql_file}. Every sheet is the output of one `-- @output` block, "
                "produced by scripts/run_queries_local.py from the four raw CSVs. "
                "Cells hold query results (values), not formulas.")
    ws["A2"].font = Font(name=FONT, italic=True, color="595959")
    ws.append([])
    ws.append(["Sheet", "Output", "Description", "Rows", "Columns"])
    for cell in ws[4]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    for sid, name, desc, n_rows, n_cols in rows:
        ws.append([sid, name, desc, n_rows, n_cols])
    for row in ws.iter_rows(min_row=5):
        for cell in row:
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[0].hyperlink = f"#'{row[0].value}'!A1"
        row[0].font = Font(name=FONT, color="0563C1", underline="single")
    for letter, width in zip("ABCDE", (12, 44, 90, 9, 9)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A5"


# ------------------------------------------------------------------- main
def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--data", default="data")
    ap.add_argument("--sql", default="sql")
    ap.add_argument("--out", default="results")
    args = ap.parse_args()

    data_dir, sql_dir, out_dir = Path(args.data), Path(args.sql), Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(":memory:")
    register_tsql_functions(con)
    print("Loading CSVs")
    load_csvs(con, data_dir)

    index_rows = []
    for sql_file in sorted(sql_dir.glob("*.sql")):
        domain = sql_file.stem
        wb = Workbook()
        wb.remove(wb.active)
        sheet_rows = []
        for name, desc, statement in split_outputs(sql_file.read_text(encoding="utf-8")):
            try:
                df = tidy(pd.read_sql_query(to_sqlite(statement), con))
            except Exception as exc:  # show which query failed, then stop
                sys.exit(f"\nFAILED {sql_file.name} :: {name}\n{exc}")
            sid = sheet_id(name)
            write_sheet(wb, sid, df)
            sheet_rows.append((sid, name, desc, len(df), len(df.columns)))
            index_rows.append((domain, sid, name, len(df), len(df.columns)))
            print(f"  {domain}.xlsx / {sid:<8} {name}  ({len(df):,} x {len(df.columns)})")
        write_index_sheet(wb, domain, sql_file.name, sheet_rows)
        wb.save(out_dir / f"{domain}.xlsx")

    lines = ["# Results index", "",
             "Generated by `scripts/run_queries_local.py` from the four raw CSVs.",
             "One workbook per domain; one sheet per `-- @output` block in the matching `sql/<domain>.sql`.", "",
             "| Workbook | Sheet | Output | Rows | Cols |",
             "|---|---|---|---:|---:|"]
    for domain, sid, name, rows, cols in index_rows:
        lines.append(f"| `{domain}.xlsx` | `{sid}` | {name} | {rows:,} | {cols} |")
    (out_dir / "_index.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"\n{len(index_rows)} sheets written to {out_dir}/")


if __name__ == "__main__":
    main()
